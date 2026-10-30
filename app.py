
import streamlit as st
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from st_aggrid import (
    AgGrid,
    GridOptionsBuilder,
    GridUpdateMode,
    DataReturnMode,
    JsCode,
)

# =========================================================
# ページ設定
# =========================================================
st.set_page_config(
    page_title="水準測量・器高式計算アプリ",
    layout="wide",
)

# =========================================================
# デザイン
# =========================================================
st.markdown(
    """
    <style>
    .stApp {
        background-color: #f1f5f9;
    }
    h1 {
        color: #1e3a5f;
    }
    div[data-testid="stNumberInput"] {
        background-color: #ffffff;
        padding: 10px;
        border-radius: 8px;
    }
    .stButton > button {
        border-radius: 6px;
    }
    </style>
    """,
    unsafe_allow_html=True,
)

# =========================================================
# タイトル・使い方
# =========================================================
st.title("水準測量・器高式計算アプリ")
st.write("水準測量で使用する器高式の計算を自動化するアプリです。")

with st.expander("使い方", expanded=True):
    st.write(
        """
        ① 基準地盤高（GH）を入力します。

        ② 水色のセルに測定値を入力します。

        ③ 黄色のセルは器高式によって自動計算されます。

        ④ 行が足りない場合は「＋ 行を追加」を押します。

        ⑤ 不要な行は選択して「− 選択行を削除」を押します。

        ⑥ 必要に応じてExcelファイルとして保存できます。
        """
    )

st.markdown("🟦 **入力**　　🟨 **自動計算**　　🟪 **累計距離**")

# =========================================================
# 基準GH
# =========================================================
base_gh = st.number_input(
    "基準地盤高（GH）",
    value=500.000,
    step=0.001,
    format="%.3f",
)

# =========================================================
# データ作成
# =========================================================
def create_empty_data(rows=10):
    return pd.DataFrame({
        "_row_id": range(rows),
        "測点": ["" for _ in range(rows)],
        "距離": [None for _ in range(rows)],
        "累計距離": [None for _ in range(rows)],
        "BS": [None for _ in range(rows)],
        "IH": [None for _ in range(rows)],
        "TP": [None for _ in range(rows)],
        "IP": [None for _ in range(rows)],
        "GH": [None for _ in range(rows)],
    })


if "survey_data" not in st.session_state:
    st.session_state.survey_data = create_empty_data()

if "grid_version" not in st.session_state:
    st.session_state.grid_version = 0


# =========================================================
# 数値変換
# =========================================================
def to_number(value):
    try:
        if pd.isna(value) or value == "":
            return None
        return float(value)
    except (ValueError, TypeError):
        return None


# =========================================================
# 器高式の自動計算
# =========================================================
def calculate_data(df, base_gh):

    df = df.copy()

    current_ih = None
    last_gh = base_gh
    total_distance = 0.0

    ih_list = []
    gh_list = []
    distance_list = []

    for _, row in df.iterrows():

        distance = to_number(row["距離"])
        bs = to_number(row["BS"])
        tp = to_number(row["TP"])
        ip = to_number(row["IP"])

        # 累計距離
        if distance is not None:
            total_distance += distance

        distance_list.append(total_distance)

        # 地盤高の計算
        gh = None

        if tp is not None and current_ih is not None:
            gh = current_ih - tp
            last_gh = gh

        elif ip is not None and current_ih is not None:
            gh = current_ih - ip

        # 後視から新しい器械高を計算
        if bs is not None:
            current_ih = last_gh + bs

        ih_list.append(current_ih)
        gh_list.append(gh)

    df["累計距離"] = distance_list
    df["IH"] = ih_list
    df["GH"] = gh_list

    return df


# =========================================================
# データ計算
# =========================================================
calculated_data = calculate_data(
    st.session_state.survey_data,
    base_gh
)

# =========================================================
# AgGridのセル色設定
# =========================================================
input_style = JsCode("""
function(params) {
    return {
        backgroundColor: '#dbeafe',
        color: '#1e293b'
    };
}
""")

auto_style = JsCode("""
function(params) {
    return {
        backgroundColor: '#fef3c7',
        color: '#1e293b'
    };
}
""")

distance_style = JsCode("""
function(params) {
    return {
        backgroundColor: '#ede9fe',
        color: '#1e293b'
    };
}
""")

# =========================================================
# 表の設定
# =========================================================
gb = GridOptionsBuilder.from_dataframe(calculated_data)

gb.configure_default_column(
    editable=False,
    resizable=True,
    sortable=False,
    filter=False,
    minWidth=90,
)

# 入力セル
for col in ["測点", "距離", "BS", "TP", "IP"]:
    gb.configure_column(
        col,
        editable=True,
        cellStyle=input_style,
    )

# 自動計算セル
for col in ["IH", "GH"]:
    gb.configure_column(
        col,
        editable=False,
        cellStyle=auto_style,
        type=["numericColumn"],
        valueFormatter="value == null ? '' : Number(value).toFixed(3)",
    )

# 累計距離
gb.configure_column(
    "累計距離",
    editable=False,
    cellStyle=distance_style,
    type=["numericColumn"],
    valueFormatter="value == null ? '' : Number(value).toFixed(3)",
)

# 内部管理用IDは非表示
gb.configure_column(
    "_row_id",
    hide=True,
)

# 行選択
gb.configure_selection(
    selection_mode="multiple",
    use_checkbox=True,
)

grid_options = gb.build()

grid_options["rowHeight"] = 42
grid_options["headerHeight"] = 45
grid_options["suppressRowClickSelection"] = True
grid_options["stopEditingWhenCellsLoseFocus"] = True

# =========================================================
# 入力表
# =========================================================
st.subheader("測量データ入力")

grid_response = AgGrid(
    calculated_data,
    gridOptions=grid_options,
    data_return_mode=DataReturnMode.AS_INPUT,
    update_mode=(
        GridUpdateMode.VALUE_CHANGED
        | GridUpdateMode.SELECTION_CHANGED
    ),
    fit_columns_on_grid_load=True,
    allow_unsafe_jscode=True,
    theme="streamlit",
    height=450,
    key=f"survey_grid_{st.session_state.grid_version}",
)

# =========================================================
# 入力内容を保存
# =========================================================
edited_data = grid_response["data"]

if not isinstance(edited_data, pd.DataFrame):
    edited_data = pd.DataFrame(edited_data)

if "_row_id" not in edited_data.columns:
    edited_data["_row_id"] = range(len(edited_data))

st.session_state.survey_data = edited_data[
    [
        "_row_id", "測点", "距離", "累計距離",
        "BS", "IH", "TP", "IP", "GH"
    ]
].copy()

# =========================================================
# 行の追加・削除
# =========================================================
col1, col2, col3 = st.columns([1, 1, 4])

with col1:
    if st.button("＋ 行を追加", use_container_width=True):

        new_data = create_empty_data(1)
        new_id = (
            int(st.session_state.survey_data["_row_id"].max()) + 1
            if len(st.session_state.survey_data) > 0
            else 0
        )

        new_data["_row_id"] = new_id

        st.session_state.survey_data = pd.concat(
            [st.session_state.survey_data, new_data],
            ignore_index=True
        )

        st.session_state.grid_version += 1
        st.rerun()

with col2:
    if st.button("− 選択行を削除", use_container_width=True):

        selected = grid_response.get("selected_rows")

        if selected is not None:

            if isinstance(selected, pd.DataFrame):
                selected_ids = selected["_row_id"].tolist()

            elif isinstance(selected, list):
                selected_ids = [
                    row["_row_id"]
                    for row in selected
                    if "_row_id" in row
                ]

            else:
                selected_ids = []

            if selected_ids:
                st.session_state.survey_data = (
                    st.session_state.survey_data[
                        ~st.session_state.survey_data["_row_id"].isin(
                            selected_ids
                        )
                    ].reset_index(drop=True)
                )

                st.session_state.grid_version += 1
                st.rerun()
            else:
                st.warning("削除する行を選択してください。")
        else:
            st.warning("削除する行を選択してください。")

# =========================================================
# Excelダウンロード
# =========================================================
st.divider()

st.subheader("Excelファイル保存")

final_data = calculate_data(
    st.session_state.survey_data,
    base_gh
)

excel_data = final_data.drop(columns=["_row_id"])

output = BytesIO()

with pd.ExcelWriter(output, engine="openpyxl") as writer:

    excel_data.to_excel(
        writer,
        index=False,
        sheet_name="水準測量野帳"
    )

    ws = writer.sheets["水準測量野帳"]

    # 基準GHを記載
    ws["K1"] = "基準地盤高（GH）"
    ws["K2"] = base_gh

    # 列幅
    for column in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column
        )
        ws.column_dimensions[
            column[0].column_letter
        ].width = max_length + 4

    # 数値の小数点表示
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if cell.column != 1:
                cell.number_format = "0.000"

st.download_button(
    label="📥 Excelファイルをダウンロード",
    data=output.getvalue(),
    file_name="水準測量野帳.xlsx",
    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    use_container_width=True,
)

st.caption("水準測量・器高式計算アプリ")
